# -*- coding: utf-8 -*-
"""구매 안내 웹 UI. 질문 → QID 매칭 → 검증 답·양식·캡처 카드. /admin 관리."""
from __future__ import annotations

import base64
import json
import math
import mimetypes
import os
import re
import threading
import time
import urllib.error
import urllib.request
from datetime import datetime, timedelta, timezone
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, unquote, urlparse

import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
WEB = Path(__file__).resolve().parent
STATIC = WEB / "static"
FORMS_DIR = WEB / "forms"
CAPTURES_DIR = WEB / "captures"
UPLOADS_DIR = WEB / "uploads"
UPLOAD_MAX_BYTES = 15 * 1024 * 1024
UPLOAD_EXTS = {
    ".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp",
    ".pdf", ".hwp", ".hwpx", ".doc", ".docx",
    ".xls", ".xlsx", ".ppt", ".pptx", ".txt", ".csv", ".zip",
}
EXCEL = ROOT / "TTA_구매챗봇_학습데이터_통합_0901.xlsx"
CACHE = ROOT / "테스트_embedding_index.json"
OVERLAY = WEB / "knowledge_overlay.json"
QUERY_LOG = WEB / "query_log.jsonl"
LOG_KEEP = 500
KST = timezone(timedelta(hours=9))
MODEL = "qwen3-embedding:0.6b"
CHAT_MODEL = "qwen3.5:4b"
GEMINI_MODELS = [
    "gemini-3.6-flash",
    "gemini-3.5-flash",
    "gemini-2.5-flash",
    "gemini-2.0-flash",
    "gemini-flash-latest",
]
QUERY_PREFIX = (
    "Instruct: Given a Korean procurement FAQ query, retrieve the matching stored question.\nQuery: "
)
REFUSE_BELOW = 0.55
AMBIGUOUS_GAP = 0.05
PORT = int(os.environ.get("PORT", "8765"))
HOST = os.environ.get("HOST", "0.0.0.0")
REFUSE_MESSAGE = (
    "자료에 없는 내용입니다. 개별 건은 추측하지 않습니다. "
    "재무인프라팀 송민규 과장(010-5110-1784)에게 문의해 주세요."
)
DISCLAIMER = "이 안내는 일반 절차입니다. 개별 건의 적격·금액·예외는 담당 확인이 필요합니다."
LOCK = threading.Lock()
GEMINI_KEY = ""
GEMINI_MODEL = ""
GEMINI_PAUSE_UNTIL = 0.0
GEMINI_PAUSE_SEC = 90
MODEL_REST: dict[str, float] = {}
MODEL_REST_QUOTA = 15 * 60
MODEL_REST_GONE = 24 * 60 * 60
MODEL_REST_BUSY = 60
EMBED_KEEP_ALIVE = "24h"
EMBED_CACHE: dict[str, list[float]] = {}
EMBED_CACHE_MAX = 500
REWRITE_CACHE: dict[str, str] = {}
REWRITE_CACHE_MAX = 400
REWRITE_STORE = WEB / "rewrite_cache.json"
QUICK_REFUSE = 0.42
CONFIDENT = 0.72
KW_SURE = 0.74
KW_GAP = 0.12
KW_MIN = 0.5
EMBED_TIMEOUT = 8
WARM_TIMEOUT = 300
WARM_EVERY = 20 * 60
WARM_TEXT = QUERY_PREFIX + "구매 절차 예열"
EMBED_READY = threading.Event()
RewarmNow = threading.Event()
WARMING_MESSAGE = (
    "안내 검색을 준비하는 중입니다. 20초쯤 뒤에 다시 물어봐 주세요. "
    "급하시면 재무인프라팀 송민규 과장(010-5110-1784)에게 문의해 주세요."
)


def load_project_env() -> dict[str, str]:
    """프로젝트 .env를 먼저 읽고, 배포용으로 GEMINI·PORT 등은 시스템 환경 변수도 받는다."""
    values: dict[str, str] = {}
    for path in (WEB / ".env", ROOT / ".env"):
        if not path.exists():
            continue
        for raw in path.read_text(encoding="utf-8").splitlines():
            line = raw.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, val = line.split("=", 1)
            values[key.strip()] = val.strip().strip('"').strip("'")
    for key in ("GEMINI_API_KEY", "GEMINI_MODEL", "PORT", "HOST"):
        if not values.get(key) and os.environ.get(key):
            values[key] = os.environ[key].strip()
    return values


def init_chat() -> None:
    global GEMINI_KEY, GEMINI_MODEL
    env = load_project_env()
    GEMINI_KEY = env.get("GEMINI_API_KEY", "").strip()
    GEMINI_MODEL = env.get("GEMINI_MODEL", "").strip()
    if not GEMINI_KEY:
        print("채팅: Ollama qwen3.5:4b (이 프로젝트 Gemini 키 없음)", flush=True)
        return
    print(f"채팅: Gemini {GEMINI_MODEL or GEMINI_MODELS[0]}", flush=True)

INDEX: list[dict] = []
META: dict[str, dict] = {}
BY_QID: dict[str, dict] = {}

EXAMPLES = [
    "참조견적 요청은 어떻게 해요?",
    "수의계약 금액 기준이 얼마야?",
    "구매요구는 어디서 하나요?",
    "경영정보에서 검수결과 등록하는 방법",
    "착수계 작성법 알려줘",
    "완료계 양식 어디 있어?",
    "잔금 줄 때 서류 뭐 필요해?",
    "전자세금계산서 받는 메일 주소",
]


def cell(value) -> str:
    if value is None or pd.isna(value):
        return ""
    text = str(value).strip()
    if not text or text == "nan":
        return ""
    return text


def split_list(value: str) -> list[str]:
    if not value:
        return []
    parts = []
    for raw in value.replace("\n", ",").split(","):
        item = raw.strip()
        if item:
            parts.append(item)
    return parts


def embed(text: str, timeout: int = EMBED_TIMEOUT) -> list[float]:
    cached = EMBED_CACHE.get(text)
    if cached:
        return cached
    body = json.dumps(
        {"model": MODEL, "input": text, "keep_alive": EMBED_KEEP_ALIVE}
    ).encode("utf-8")
    req = urllib.request.Request(
        "http://127.0.0.1:11434/api/embed",
        data=body,
        headers={"Content-Type": "application/json"},
    )
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        data = json.loads(resp.read().decode("utf-8"))
    vec = data["embeddings"][0]
    with LOCK:
        if len(EMBED_CACHE) >= EMBED_CACHE_MAX:
            EMBED_CACHE.clear()
        EMBED_CACHE[text] = vec
    return vec


def keep_warm() -> None:
    """임베딩 모델을 올려 두고 계속 붙잡는다.

    적재에 1분 넘게 걸리는 기기가 있으므로 기다리는 일은 여기서만 한다.
    질문 쪽은 준비 신호를 보고 판단하니 헛되게 기다리지 않는다.
    """
    while True:
        start = time.perf_counter()
        try:
            if not EMBED_READY.is_set():
                EMBED_CACHE.pop(WARM_TEXT, None)
            embed(WARM_TEXT, timeout=WARM_TIMEOUT)
            if not EMBED_READY.is_set():
                print(f"임베딩 준비 완료 ({time.perf_counter() - start:.1f}초)", flush=True)
            EMBED_READY.set()
        except Exception as e:
            EMBED_READY.clear()
            EMBED_CACHE.pop(WARM_TEXT, None)
            print(f"임베딩 준비 실패: {e}", flush=True)
        wait = WARM_EVERY if EMBED_READY.is_set() else 5
        RewarmNow.clear()
        RewarmNow.wait(timeout=wait)


def strip_think(text: str) -> str:
    text = re.sub(r"<think>.*?</think>", "", text, flags=re.S | re.I)
    return text.strip()


def gemini_ready() -> bool:
    return bool(GEMINI_KEY) and time.time() >= GEMINI_PAUSE_UNTIL


def pause_gemini(reason: str) -> None:
    """모두 실패하면 잠시 쉰다. 매 질문마다 헛된 호출을 반복하지 않는다."""
    global GEMINI_PAUSE_UNTIL
    if time.time() < GEMINI_PAUSE_UNTIL:
        return
    GEMINI_PAUSE_UNTIL = time.time() + GEMINI_PAUSE_SEC
    print(f"Gemini {GEMINI_PAUSE_SEC}초 쉬는 중: {reason}", flush=True)


def rest_model(model: str, seconds: float, reason: str) -> None:
    """모델별로 쉬게 둔다. 하루 한도를 쓴 모델을 계속 두드리지 않는다."""
    MODEL_REST[model] = time.time() + seconds
    print(f"Gemini {model} 쉬는 중({int(seconds)}초): {reason}", flush=True)


def gemini_models_to_try() -> list[str]:
    now = time.time()
    order = ([GEMINI_MODEL] if GEMINI_MODEL else []) + [
        m for m in GEMINI_MODELS if m != GEMINI_MODEL
    ]
    live = [m for m in order if MODEL_REST.get(m, 0) <= now]
    return live or order[:1]


def gemini_call(model: str, system: str, user: str, num_predict: int, thinking: bool) -> dict:
    gen_cfg = {"temperature": 0.1, "maxOutputTokens": max(256, num_predict)}
    if thinking:
        gen_cfg["thinkingConfig"] = {"thinkingBudget": 0}
    payload = {
        "systemInstruction": {"parts": [{"text": system}]},
        "contents": [{"role": "user", "parts": [{"text": user}]}],
        "generationConfig": gen_cfg,
    }
    req = urllib.request.Request(
        "https://generativelanguage.googleapis.com/v1beta/models/"
        f"{model}:generateContent",
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json", "x-goog-api-key": GEMINI_KEY},
    )
    with urllib.request.urlopen(req, timeout=60) as resp:
        return json.loads(resp.read().decode("utf-8"))


def gemini_complete(system: str, user: str, num_predict: int = 500) -> str:
    global GEMINI_MODEL
    if not GEMINI_KEY:
        raise RuntimeError("GEMINI_API_KEY가 없습니다")
    last_err = ""
    for model in gemini_models_to_try():
        for thinking in (True, False):
            try:
                data = gemini_call(model, system, user, num_predict, thinking)
            except urllib.error.HTTPError as e:
                last_err = f"HTTP {e.code} " + e.read().decode("utf-8", errors="replace")[:200]
                if e.code in {429, 403}:
                    rest_model(model, MODEL_REST_QUOTA, f"HTTP {e.code} 한도")
                elif e.code == 404:
                    rest_model(model, MODEL_REST_GONE, "이 키로 쓸 수 없는 모델")
                elif e.code >= 500:
                    rest_model(model, MODEL_REST_BUSY, f"HTTP {e.code}")
                else:
                    continue  # 400은 thinkingConfig 없이 한 번 더
                break
            except (urllib.error.URLError, TimeoutError) as e:
                last_err = f"연결 실패: {e}"
                rest_model(model, MODEL_REST_BUSY, str(e))
                break
            cand = (data.get("candidates") or [{}])[0]
            parts = (cand.get("content") or {}).get("parts") or []
            text = "".join(
                p.get("text") or "" for p in parts if not p.get("thought")
            ).strip()
            if text:
                if GEMINI_MODEL != model:
                    GEMINI_MODEL = model
                    print(f"채팅 모델: Gemini {model}", flush=True)
                return text
            last_err = json.dumps(data, ensure_ascii=False)[:200]
    pause_gemini(last_err or "모든 모델 실패")
    raise RuntimeError(last_err or "Gemini 응답이 비었습니다")


def chat_complete(system: str, user: str, num_predict: int = 500) -> str:
    if GEMINI_KEY:
        if not gemini_ready():
            raise RuntimeError("Gemini 잠시 쉬는 중")
        return strip_think(gemini_complete(system, user, num_predict))
    body = json.dumps(
        {
            "model": CHAT_MODEL,
            "stream": False,
            "think": False,
            "options": {"temperature": 0.1, "num_predict": num_predict},
            "messages": [
                {"role": "system", "content": system},
                {"role": "user", "content": user},
            ],
        }
    ).encode("utf-8")
    req = urllib.request.Request(
        "http://127.0.0.1:11434/api/chat",
        data=body,
        headers={"Content-Type": "application/json"},
    )
    with urllib.request.urlopen(req, timeout=180) as resp:
        data = json.loads(resp.read().decode("utf-8"))
    return strip_think((data.get("message") or {}).get("content") or "")


ROUTE_SYSTEM = """당신은 TTA 구매 안내의 앞단이다. 직원 질문과 바로 앞 대화를 이해한다.
JSON만 출력한다. 설명 문장은 쓰지 않는다.

구매·발주·계약·검수·대금·세금계산서·입찰·하자보증·회계지출 등 구매 업무 절차 질문이면:
{"action":"search","query":"저장된 FAQ에서 찾을 한국어 한 문장"}

바로 앞 안내가 있고, 그/그거/그 양식/그 화면/이어서/그럼처럼 그 안내를 가리키면:
{"action":"same"}

구매 업무가 아니거나(날씨·잡담 등), 특정 업체·특정 금액의 가부·적격·예외 승인을 묻는 질문이면 앞 안내가 있어도:
{"action":"refuse"}"""

PICK_SYSTEM = """역할: TTA 구매 안내 후보 고르기.
JSON만 출력한다. 설명 문장은 쓰지 않는다.

직원 질문에 가장 잘 맞는 안내 번호만 고른다.
목록에 없는 번호를 만들지 않는다.
안내 내용을 추측하거나 답을 쓰지 않는다.

하나만 분명하면:
{"action":"pick","qid":"PO-001"}

둘 이상이 비슷하면:
{"action":"clarify","qids":["PO-001","CT-002"]}

하나도 해당하지 않으면:
{"action":"refuse"}"""

REWRITE_SYSTEM = """역할: TTA 구매 안내 문장을 직원에게 읽기 쉽게 다시 쓰는 담당.
한국어만 쓴다. 말투는 짧고 절차 중심이다.
직원이 지금 물은 부분에 먼저 답한다.

절대 규칙:
- 아래에 준 검증된 답 안에 있는 내용만 쓴다.
- 금액 기준, 양식 이름, 메뉴 경로, 서류 목록을 지어내지 않는다.
- 검증된 답에 없는 규정·예외·판단은 쓰지 않는다.
- 개별 건의 가부·적격을 판단하지 않는다.
- 양식은 관련 양식 칸에 있는 이름만 언급한다. 없으면 양식을 말하지 않는다.
- 없는 내용을 보완하려고 일반 지식을 섞지 않는다."""

SEARCH_K = 8
PICK_MIN = 0.50
CLARIFY_MESSAGE = "가까운 안내가 여러 건입니다. 해당하는 항목을 골라 주세요."
FOLLOW_HINT = re.compile(
    r"(그\s*(양식|화면|절차|건|거|것|안내)|그거|이어서|그럼|그러면|위에|방금|앞에서)"
)


def parse_json_obj(raw: str) -> dict:
    match = re.search(r"\{.*\}", raw or "", flags=re.S)
    if not match:
        return {}
    try:
        data = json.loads(match.group(0))
    except Exception:
        return {}
    return data if isinstance(data, dict) else {}


def format_history(history: list | None) -> str:
    if not history:
        return "(이전 대화 없음)"
    lines = []
    for item in history[-8:]:
        if not isinstance(item, dict):
            continue
        if item.get("role") == "user":
            text = str(item.get("text") or "").strip()[:240]
            if text:
                lines.append(f"직원: {text}")
            continue
        if item.get("refused"):
            lines.append("안내: 자료 없음으로 거절")
            continue
        qid = str(item.get("qid") or "").strip()
        question = str(item.get("question") or "").strip()[:200]
        if qid or question:
            lines.append(f"안내: [{qid}] {question}".strip())
    return "\n".join(lines) or "(이전 대화 없음)"


def last_qid(history: list | None) -> str:
    for item in reversed(history or []):
        if not isinstance(item, dict):
            continue
        if item.get("role") == "user" or item.get("refused"):
            continue
        qid = str(item.get("qid") or "").strip()
        if qid:
            return qid
    return ""


def route_question(query: str, history: list | None = None) -> dict:
    user = f"이전 대화:\n{format_history(history)}\n\n지금 질문:\n{query}"
    try:
        data = parse_json_obj(chat_complete(ROUTE_SYSTEM, user, num_predict=120))
        action = str(data.get("action") or "").strip()
        if action == "refuse":
            return {"action": "refuse"}
        if action == "same":
            return {"action": "same"}
        rewritten = str(data.get("query") or "").strip()
        return {"action": "search", "query": rewritten or query}
    except Exception:
        return {"action": "search", "query": query}


def pick_guide(query: str, ranked: list[tuple[float, dict]], history: list | None = None) -> dict:
    lines = []
    allowed = set()
    for score, row in ranked:
        qid = row["qid"]
        allowed.add(qid)
        lines.append(
            f"{qid} | {row.get('cat') or ''} | {row['q']} | {score:.3f}"
        )
    user = (
        f"이전 대화:\n{format_history(history)}\n\n"
        f"지금 질문:\n{query}\n\n"
        f"후보(안내번호 | 분류 | 저장된 질문 | 비슷함):\n"
        + "\n".join(lines)
    )
    try:
        data = parse_json_obj(chat_complete(PICK_SYSTEM, user, num_predict=160))
    except Exception:
        return {"action": "fallback"}
    action = str(data.get("action") or "").strip()
    if action == "refuse":
        return {"action": "refuse"}
    if action == "clarify":
        qids = []
        for qid in data.get("qids") or []:
            qid = str(qid).strip()
            if qid in allowed and qid not in qids:
                qids.append(qid)
        if len(qids) >= 2:
            return {"action": "clarify", "qids": qids[:4]}
        if len(qids) == 1:
            return {"action": "pick", "qid": qids[0]}
        return {"action": "fallback"}
    qid = str(data.get("qid") or "").strip()
    if action == "pick" and qid in allowed:
        return {"action": "pick", "qid": qid}
    return {"action": "fallback"}


def ranked_by_qid(ranked: list[tuple[float, dict]], qid: str) -> tuple[float, dict] | None:
    for score, row in ranked:
        if row["qid"] == qid:
            return score, row
    return None


def refuse_result(score: float | None = None, ranked: list[tuple[float, dict]] | None = None) -> dict:
    candidates = []
    if ranked:
        candidates = [
            {
                "qid": r["qid"],
                "category": r.get("cat", ""),
                "question": r["q"],
                "score": round(s, 3),
            }
            for s, r in ranked[:3]
        ]
    return {
        "refused": True,
        "score": None if score is None else round(score, 3),
        "message": REFUSE_MESSAGE,
        "candidates": candidates,
        "disclaimer": DISCLAIMER,
    }


def clarify_result(ranked: list[tuple[float, dict]], qids: list[str]) -> dict:
    wanted = set(qids)
    candidates = [
        {
            "qid": r["qid"],
            "category": r.get("cat", ""),
            "question": r["q"],
            "score": round(s, 3),
        }
        for s, r in ranked
        if r["qid"] in wanted
    ]
    return {
        "refused": False,
        "clarify": True,
        "ambiguous": True,
        "score": candidates[0]["score"] if candidates else None,
        "message": CLARIFY_MESSAGE,
        "candidates": candidates,
        "disclaimer": DISCLAIMER,
    }


def load_rewrite_cache() -> None:
    if not REWRITE_STORE.exists():
        return
    try:
        data = json.loads(REWRITE_STORE.read_text(encoding="utf-8"))
    except Exception:
        return
    if isinstance(data, dict):
        REWRITE_CACHE.update({str(k): str(v) for k, v in data.items() if v})


def rewrite_key(qid: str, query: str) -> str:
    norm = re.sub(r"\s+", " ", (query or "").strip().lower())
    return f"{qid}|{norm}"


def remember_rewrite(key: str, text: str) -> None:
    with LOCK:
        if len(REWRITE_CACHE) >= REWRITE_CACHE_MAX:
            REWRITE_CACHE.clear()
        REWRITE_CACHE[key] = text
        try:
            REWRITE_STORE.write_text(
                json.dumps(REWRITE_CACHE, ensure_ascii=False), encoding="utf-8"
            )
        except Exception:
            pass


def rewrite_answer(staff_query: str, packed: dict) -> str:
    source = (packed.get("answer") or "").strip()
    if not source:
        return ""
    key = rewrite_key(packed.get("qid") or "", staff_query or packed.get("question") or "")
    hit = REWRITE_CACHE.get(key)
    if hit:
        return hit
    forms = ", ".join(f.get("name") for f in (packed.get("forms") or []) if f.get("name"))
    user = (
        f"직원 질문: {staff_query or packed.get('question') or ''}\n"
        f"저장된 질문: {packed.get('question') or ''}\n"
        f"분류: {packed.get('category') or ''}\n"
        f"관련 양식: {forms or '없음'}\n\n"
        f"검증된 답:\n{source}\n\n"
        "직원이 지금 물은 부분에 먼저 답하되, 위 검증된 답만 사용하세요. "
        "양식·파일을 물으면 관련 양식 이름을 먼저 적고, 그 밖에 양식 설명을 지어내지 마세요."
    )
    try:
        text = chat_complete(REWRITE_SYSTEM, user, num_predict=900)
    except Exception:
        return source
    if text:
        remember_rewrite(key, text)
    return text or source


def form_files(names: list[str]) -> list[dict]:
    out = []
    for name in names:
        stem = Path(name).name
        found = None
        if FORMS_DIR.exists():
            if Path(stem).suffix and (FORMS_DIR / stem).exists():
                found = stem
            else:
                for ext in (".pdf", ".hwp", ".hwpx", ".xlsx", ".xls", ".docx"):
                    candidate = FORMS_DIR / f"{stem}{ext}"
                    if candidate.exists():
                        found = candidate.name
                        break
        out.append({
            "name": name,
            "file": found,
            "url": f"/forms/{found}" if found else None,
        })
    return out


def capture_steps(qid: str, notes: list[str]) -> list[dict]:
    folder = CAPTURES_DIR / qid
    steps = []
    if folder.exists():
        images = sorted(
            p for p in folder.iterdir()
            if p.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp"}
        )
        for i, path in enumerate(images):
            title = notes[i] if i < len(notes) else path.stem
            steps.append({
                "step": i + 1,
                "title": title,
                "file": path.name,
                "url": f"/captures/{qid}/{path.name}",
            })
        return steps
    for i, title in enumerate(notes[:6]):
        steps.append({"step": i + 1, "title": title, "file": None, "url": None})
    return steps


def load_overlay() -> dict:
    if not OVERLAY.exists():
        return {}
    try:
        return json.loads(OVERLAY.read_text(encoding="utf-8"))
    except Exception:
        return {}


def write_overlay(data: dict) -> None:
    OVERLAY.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def load_meta_from_excel() -> dict[str, dict]:
    df = pd.read_excel(EXCEL, sheet_name=0, header=2, dtype=str)
    out: dict[str, dict] = {}
    for _, r in df.iterrows():
        qid = cell(r.iloc[0])
        if not qid or qid.startswith("▶"):
            continue
        out[qid] = {
            "answer": cell(r.iloc[3]),
            "forms": split_list(cell(r.iloc[6])),
            "similar": split_list(cell(r.iloc[7])),
            "capture_notes": split_list(cell(r.iloc[12])),
            "rule": cell(r.iloc[13]),
            "law": cell(r.iloc[14]) if len(r) > 14 else "",
        }
    return out


def apply_overlay(meta: dict[str, dict]) -> dict[str, dict]:
    for qid, extra in load_overlay().items():
        base = meta.setdefault(
            qid,
            {"answer": "", "forms": [], "similar": [], "capture_notes": [], "rule": "", "law": ""},
        )
        for key in ("answer", "rule", "law"):
            if key in extra and extra[key] is not None:
                base[key] = extra[key]
        for key in ("forms", "similar", "capture_notes"):
            if key in extra and extra[key] is not None:
                base[key] = extra[key]
    return meta


def load_meta() -> dict[str, dict]:
    return apply_overlay(load_meta_from_excel())


def reload_knowledge() -> None:
    """관리자가 답을 고치면 다듬은 문장 캐시도 버린다. 옛 답이 남지 않게 한다."""
    global META
    META = load_meta()
    with LOCK:
        REWRITE_CACHE.clear()
        try:
            REWRITE_STORE.unlink(missing_ok=True)
        except Exception:
            pass


def load_index() -> list[dict]:
    if not CACHE.exists():
        raise FileNotFoundError(
            "색인 파일이 없습니다. 먼저 python 테스트_embedding.py 를 한 번 실행하세요."
        )
    data = json.loads(CACHE.read_text(encoding="utf-8"))
    rows = data["rows"]
    for row in rows:
        row["norm"] = math.sqrt(sum(x * x for x in row["vec"])) or 1.0
        row["bi"] = bigrams(row["q"])
    return rows


def bigrams(text: str) -> set[str]:
    clean = re.sub(r"[^0-9a-z가-힣]", "", (text or "").lower())
    if len(clean) < 2:
        return {clean} if clean else set()
    return {clean[i:i + 2] for i in range(len(clean) - 1)}


def keyword_rank(query: str, k: int = 3) -> list[tuple[float, dict]]:
    """글자 겹침만으로 순위를 낸다. 임베딩이 느려도 즉시 답한다."""
    qb = bigrams(query)
    if not qb:
        return []
    scored = [
        (len(qb & (row.get("bi") or set())) / len(qb), row) for row in INDEX
    ]
    scored.sort(key=lambda x: x[0], reverse=True)
    return scored[:k]


def vector_rank(query: str, k: int = 3) -> list[tuple[float, dict]]:
    qv = embed(QUERY_PREFIX + query)
    qn = math.sqrt(sum(x * x for x in qv)) or 1.0
    scored = []
    for row in INDEX:
        dot = sum(x * y for x, y in zip(qv, row["vec"]))
        scored.append((dot / (qn * row["norm"]), row))
    scored.sort(key=lambda x: x[0], reverse=True)
    return scored[:k]


def keyword_fallback(query: str, k: int) -> list[tuple[float, dict]]:
    out = []
    for cover, row in keyword_rank(query, k):
        if cover < KW_MIN:
            continue
        out.append((min(0.95, 0.5 + cover / 2), row))
    return out


def search(query: str, k: int = 3) -> list[tuple[float, dict]]:
    if not EMBED_READY.is_set():
        # 모델 적재를 기다리며 시간을 버리지 않는다. 글자 겹침으로 찾아 본다.
        RewarmNow.set()
        return keyword_fallback(query, k)
    try:
        return vector_rank(query, k)
    except Exception as e:
        print(f"임베딩 실패 → 글자 검색으로 대체: {e}", flush=True)
        EMBED_READY.clear()
        RewarmNow.set()
        return keyword_fallback(query, k)


def pack_row(row: dict, score: float | None, ranked: list[tuple[float, dict]] | None) -> dict:
    extra = META.get(row["qid"], {})
    candidates = []
    if ranked:
        candidates = [
            {
                "qid": r["qid"],
                "category": r.get("cat", ""),
                "question": r["q"],
                "score": round(s, 3),
            }
            for s, r in ranked
        ]
    return {
        "refused": False,
        "score": None if score is None else round(score, 3),
        "qid": row["qid"],
        "category": row.get("cat", ""),
        "question": row["q"],
        "answer": extra.get("answer") or "",
        "forms": form_files(extra.get("forms") or []),
        "captures": capture_steps(row["qid"], extra.get("capture_notes") or []),
        "similar": extra.get("similar") or [],
        "rule": extra.get("rule") or "",
        "law": extra.get("law") or "",
        "disclaimer": DISCLAIMER,
        "candidates": candidates,
        "ambiguous": False,
    }


def finish(result: dict, query: str) -> dict:
    """검증된 답을 AI가 다듬을 때까지 기다린 뒤 한 번에 돌려준다."""
    if not (result.get("answer") or "").strip():
        return result
    result["answer"] = rewrite_answer(query, result)
    return result


def decide_locally(ranked: list[tuple[float, dict]], query: str) -> dict:
    """검색 점수만으로 결론이 나는 경우를 골라낸다. LLM 왕복을 줄인다."""
    if not ranked:
        return {"action": "refuse"}
    top_s = ranked[0][0]
    gap = top_s - ranked[1][0] if len(ranked) >= 2 else 1.0
    if top_s < QUICK_REFUSE:
        return {"action": "refuse"}
    if top_s >= REFUSE_BELOW and gap >= AMBIGUOUS_GAP:
        # 확신이 낮으면 답과 함께 가까운 안내도 같이 보여 준다.
        return {"action": "top", "show_others": top_s < CONFIDENT}
    if top_s >= REFUSE_BELOW:
        return {"action": "pick"}
    return {"action": "route"}


def ask(query: str, history: list | None = None) -> dict:
    history = [h for h in (history or []) if isinstance(h, dict)][:8]
    prior = last_qid(history)

    # 이어지는 질문이면 앞 안내를 그대로 쓴다. LLM을 부르지 않는다.
    if prior and FOLLOW_HINT.search(query):
        row = BY_QID.get(prior)
        if row:
            return finish(pack_row(row, None, None), query)

    # 저장된 질문과 글자가 거의 같으면 임베딩 없이 바로 답한다.
    kw = keyword_rank(query, k=SEARCH_K)
    if kw and kw[0][0] >= KW_SURE:
        gap = kw[0][0] - kw[1][0] if len(kw) >= 2 else 1.0
        if gap >= KW_GAP:
            cover, row = kw[0]
            return finish(pack_row(row, round(cover, 3), None), query)

    ranked = search(query, k=SEARCH_K)
    if not ranked and not EMBED_READY.is_set():
        return {
            "refused": True,
            "warming": True,
            "message": WARMING_MESSAGE,
            "candidates": [],
            "disclaimer": DISCLAIMER,
        }
    plan = decide_locally(ranked, query)

    # 확실히 맞거나 확실히 무관하면 여기서 끝낸다.
    if plan["action"] == "refuse" and not ranked:
        return refuse_result()
    if plan["action"] == "top":
        top_s, top = ranked[0]
        result = pack_row(top, top_s, ranked[:3])
        result["ambiguous"] = bool(plan.get("show_others"))
        return finish(result, query)
    if plan["action"] == "refuse":
        return refuse_result(ranked[0][0], ranked)

    # 후보가 가깝거나 점수가 애매해도 첫 화면은 기다리지 않는다.
    # 가장 가까운 안내를 바로 주고, 다른 후보는 같이 보여 고르게 한다.
    top_s, top = ranked[0]
    if top_s < REFUSE_BELOW:
        return refuse_result(top_s, ranked)
    result = pack_row(top, top_s, ranked[:3])
    result["ambiguous"] = True
    return finish(result, query)


def guide(qid: str, query: str = "") -> dict:
    row = BY_QID.get(qid)
    if not row:
        return {"refused": True, "message": REFUSE_MESSAGE, "disclaimer": DISCLAIMER}
    return finish(pack_row(row, None, None), query)


def polish(qid: str, query: str = "") -> dict:
    row = BY_QID.get(qid)
    if not row:
        return {"answer": ""}
    packed = pack_row(row, None, None)
    return {"answer": rewrite_answer(query, packed), "qid": qid}


def safe_upload_name(name: str) -> str:
    stem = Path(str(name or "")).name
    stem = re.sub(r"[^\w.\- ()가-힣]", "_", stem).strip() or "첨부파일"
    return stem[:120]


def save_upload(name: str, data_b64: str) -> dict:
    clean = safe_upload_name(name)
    ext = Path(clean).suffix.lower()
    if ext not in UPLOAD_EXTS:
        raise ValueError(f"올릴 수 없는 형식입니다: {ext or '확장자 없음'}")
    try:
        blob = base64.b64decode((data_b64 or "").split(",", 1)[-1], validate=False)
    except Exception as exc:
        raise ValueError("파일을 읽지 못했습니다") from exc
    if not blob:
        raise ValueError("빈 파일입니다")
    if len(blob) > UPLOAD_MAX_BYTES:
        raise ValueError("파일이 너무 큽니다. 15MB까지 올릴 수 있습니다")
    day = datetime.now(KST).strftime("%Y%m%d")
    folder = UPLOADS_DIR / day
    folder.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(KST).strftime("%H%M%S")
    target = folder / f"{stamp}_{clean}"
    n = 1
    while target.exists():
        target = folder / f"{stamp}_{n}_{clean}"
        n += 1
    target.write_bytes(blob)
    return {
        "name": clean,
        "url": f"/uploads/{day}/{target.name}",
        "size": len(blob),
        "image": ext in {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"},
    }


def append_query_log(
    user: str, query: str, result: dict, source: str, files: list | None = None
) -> None:
    user = (user or "").strip()[:40] or "이름 없음"
    query = (query or "").strip()[:500]
    answer = result.get("answer") or result.get("message") or ""
    entry = {
        "time": datetime.now(KST).strftime("%Y-%m-%d %H:%M:%S"),
        "user": user,
        "query": query,
        "source": source,
        "refused": bool(result.get("refused")),
        "qid": result.get("qid") or "",
        "matched": result.get("question") or "",
        "score": result.get("score"),
        "ambiguous": bool(result.get("ambiguous")),
        "answer": answer[:2000],
        "forms": [f.get("name") for f in (result.get("forms") or []) if f.get("name")],
        "files": [
            {"name": f.get("name"), "url": f.get("url")}
            for f in (files or [])
            if isinstance(f, dict) and f.get("url")
        ],
    }
    with LOCK:
        lines = QUERY_LOG.read_text(encoding="utf-8").splitlines() if QUERY_LOG.exists() else []
        lines.append(json.dumps(entry, ensure_ascii=False))
        QUERY_LOG.write_text("\n".join(lines[-LOG_KEEP:]) + "\n", encoding="utf-8")


def read_query_log(limit: int = 200) -> list[dict]:
    if not QUERY_LOG.exists():
        return []
    rows = []
    for line in QUERY_LOG.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        try:
            rows.append(json.loads(line))
        except Exception:
            continue
    rows.reverse()
    return rows[:limit]


def admin_catalog() -> dict:
    cats = []
    seen = set()
    items = []
    answered = 0
    form_ready = 0
    cap_ready = 0
    for row in INDEX:
        extra = META.get(row["qid"], {})
        forms = form_files(extra.get("forms") or [])
        caps = [c for c in capture_steps(row["qid"], extra.get("capture_notes") or []) if c.get("url")]
        has_answer = bool(extra.get("answer"))
        if has_answer:
            answered += 1
        if any(f.get("url") for f in forms):
            form_ready += 1
        if caps:
            cap_ready += 1
        cat = row.get("cat") or ""
        if cat and cat not in seen:
            seen.add(cat)
            cats.append(cat)
        items.append({
            "qid": row["qid"],
            "category": cat,
            "question": row["q"],
            "has_answer": has_answer,
            "forms_ready": sum(1 for f in forms if f.get("url")),
            "forms_total": len(forms),
            "captures": len(caps),
        })
    return {
        "indexed": len(INDEX),
        "answered": answered,
        "form_ready": form_ready,
        "cap_ready": cap_ready,
        "categories": cats,
        "items": items,
    }


def admin_item(qid: str) -> dict | None:
    row = BY_QID.get(qid)
    if not row:
        return None
    extra = META.get(qid, {})
    return {
        "qid": qid,
        "category": row.get("cat", ""),
        "question": row["q"],
        "answer": extra.get("answer") or "",
        "forms": form_files(extra.get("forms") or []),
        "similar": extra.get("similar") or [],
        "capture_notes": extra.get("capture_notes") or [],
        "captures": capture_steps(qid, extra.get("capture_notes") or []),
        "rule": extra.get("rule") or "",
        "law": extra.get("law") or "",
    }


def try_write_excel(qid: str, fields: dict) -> str:
    try:
        from openpyxl import load_workbook
        wb = load_workbook(EXCEL)
        ws = wb.worksheets[0]
        target = None
        for row in range(4, ws.max_row + 1):
            if cell(ws.cell(row, 1).value) == qid:
                target = row
                break
        if not target:
            wb.close()
            return "엑셀에서 해당 QID를 찾지 못했습니다."
        ws.cell(target, 4).value = fields.get("answer") or None
        ws.cell(target, 7).value = ", ".join(fields.get("forms") or []) or None
        ws.cell(target, 8).value = ", ".join(fields.get("similar") or []) or None
        ws.cell(target, 13).value = ", ".join(fields.get("capture_notes") or []) or None
        wb.save(EXCEL)
        wb.close()
        return ""
    except PermissionError:
        return "엑셀이 다른 프로그램에서 열려 있어 엑셀 파일은 수정하지 못했습니다. 관리자 저장본에는 반영됐습니다."
    except Exception as e:
        return f"엑셀 저장 실패: {e}"


def save_item(data: dict) -> dict:
    qid = str(data.get("qid", "")).strip()
    if qid not in BY_QID:
        return {"error": "없는 안내 번호입니다"}
    answer = str(data.get("answer", "")).strip()
    forms = [x.strip() for x in (data.get("forms") or []) if str(x).strip()]
    similar = [x.strip() for x in (data.get("similar") or []) if str(x).strip()]
    notes = [x.strip() for x in (data.get("capture_notes") or []) if str(x).strip()]
    fields = {"answer": answer, "forms": forms, "similar": similar, "capture_notes": notes}
    with LOCK:
        overlay = load_overlay()
        overlay[qid] = {**overlay.get(qid, {}), **fields}
        write_overlay(overlay)
        excel_msg = try_write_excel(qid, fields)
        reload_knowledge()
    return {"ok": True, "excel_message": excel_msg, "item": admin_item(qid)}


def safe_filename(name: str) -> str:
    name = Path(name).name
    if not name or name in {".", ".."}:
        raise ValueError("파일 이름이 없습니다")
    if re.search(r'[<>:"/\\|?*]', name):
        raise ValueError("파일 이름에 쓸 수 없는 문자가 있습니다")
    return name


def decode_upload(data: dict) -> tuple[str, bytes]:
    name = safe_filename(str(data.get("filename", "")))
    raw = base64.b64decode(str(data.get("data", "")), validate=False)
    if not raw:
        raise ValueError("파일이 비어 있습니다")
    return name, raw


def upload_form(data: dict) -> dict:
    name, raw = decode_upload(data)
    stem = str(data.get("form_name") or Path(name).stem).strip()
    ext = Path(name).suffix.lower() or ".pdf"
    if ext not in {".pdf", ".hwp", ".hwpx", ".xlsx", ".xls", ".docx"}:
        return {"error": "올릴 수 있는 양식은 pdf, hwp, xlsx, docx 입니다"}
    dest = FORMS_DIR / f"{stem}{ext}"
    dest.write_bytes(raw)
    reload_knowledge()
    return {"ok": True, "file": dest.name, "url": f"/forms/{dest.name}"}


def upload_capture(data: dict) -> dict:
    qid = str(data.get("qid", "")).strip()
    if qid not in BY_QID:
        return {"error": "없는 안내 번호입니다"}
    name, raw = decode_upload(data)
    ext = Path(name).suffix.lower()
    if ext not in {".png", ".jpg", ".jpeg", ".webp"}:
        return {"error": "이미지는 png, jpg, webp 만 됩니다"}
    folder = CAPTURES_DIR / qid
    folder.mkdir(parents=True, exist_ok=True)
    dest = folder / name
    dest.write_bytes(raw)
    reload_knowledge()
    return {"ok": True, "item": admin_item(qid)}


def delete_capture(data: dict) -> dict:
    qid = str(data.get("qid", "")).strip()
    name = safe_filename(str(data.get("filename", "")))
    path = CAPTURES_DIR / qid / name
    if not path.exists() or not path.is_file():
        return {"error": "파일이 없습니다"}
    path.unlink()
    reload_knowledge()
    return {"ok": True, "item": admin_item(qid)}


class Handler(BaseHTTPRequestHandler):
    def log_message(self, fmt: str, *args) -> None:
        print("[web]", fmt % args)

    def _send(self, code: int, body: bytes, content_type: str) -> None:
        self.send_response(code)
        self.send_header("Content-Type", content_type)
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(body)

    def _json(self, code: int, payload: dict) -> None:
        self._send(code, json.dumps(payload, ensure_ascii=False).encode("utf-8"), "application/json; charset=utf-8")

    def _file(self, path: Path) -> None:
        if not path.exists() or not path.is_file():
            self._send(404, b"not found", "text/plain")
            return
        ctype = mimetypes.guess_type(str(path))[0] or "application/octet-stream"
        self._send(200, path.read_bytes(), ctype)

    def do_GET(self) -> None:
        parsed = urlparse(self.path)
        path = unquote(parsed.path)
        if path in ("/", "/index.html"):
            self._file(STATIC / "index.html")
            return
        if path in ("/admin", "/admin.html"):
            self._file(STATIC / "admin.html")
            return
        if path == "/api/health":
            self._json(200, {"ok": True, "indexed": len(INDEX), "model": MODEL})
            return
        if path == "/api/examples":
            self._json(200, {"examples": EXAMPLES})
            return
        if path == "/api/admin/catalog":
            self._json(200, admin_catalog())
            return
        if path == "/api/admin/item":
            qid = parse_qs(parsed.query).get("qid", [""])[0]
            item = admin_item(qid)
            if not item:
                self._json(404, {"error": "없는 안내 번호입니다"})
                return
            self._json(200, item)
            return
        if path == "/api/admin/logs":
            self._json(200, {"logs": read_query_log(200)})
            return
        if path.startswith("/forms/"):
            name = Path(path[len("/forms/"):]).name
            self._file(FORMS_DIR / name)
            return
        if path.startswith("/captures/"):
            rel = Path(path[len("/captures/"):])
            if ".." in rel.parts:
                self._send(400, b"bad path", "text/plain")
                return
            self._file(CAPTURES_DIR / rel)
            return
        if path.startswith("/uploads/"):
            rel = Path(path[len("/uploads/"):])
            if ".." in rel.parts:
                self._send(400, b"bad path", "text/plain")
                return
            self._file(UPLOADS_DIR / rel)
            return
        self._send(404, b"not found", "text/plain")

    def do_POST(self) -> None:
        path = urlparse(self.path).path
        length = int(self.headers.get("Content-Length", "0"))
        raw = self.rfile.read(length)
        try:
            data = json.loads(raw.decode("utf-8")) if raw else {}
        except Exception:
            self._json(400, {"error": "잘못된 요청"})
            return

        if path == "/api/upload":
            items = data.get("files")
            if not isinstance(items, list) or not items:
                self._json(400, {"error": "올릴 파일이 없습니다"})
                return
            saved = []
            try:
                for item in items[:5]:
                    saved.append(save_upload(item.get("name"), item.get("data")))
            except ValueError as e:
                self._json(400, {"error": str(e)})
                return
            except Exception as e:
                self._json(500, {"error": str(e)})
                return
            self._json(200, {"files": saved})
            return

        if path == "/api/ask":
            query = str(data.get("query", "")).strip()
            user = str(data.get("user", "")).strip()
            files = data.get("files") if isinstance(data.get("files"), list) else []
            if not query:
                self._json(400, {"error": "질문을 입력하세요"})
                return
            try:
                history = data.get("history") if isinstance(data.get("history"), list) else []
                started = time.perf_counter()
                result = ask(query, history)
                print(f"[답변] {time.perf_counter() - started:.1f}초 · {query[:30]}", flush=True)
                if files:
                    result["files"] = files
                append_query_log(user, query, result, "질문", files)
                self._json(200, result)
            except urllib.error.URLError as e:
                self._json(502, {"error": "Ollama에 연결하지 못했습니다. Ollama가 실행 중인지 확인하세요.", "detail": str(e)})
            except Exception as e:
                self._json(500, {"error": str(e)})
            return

        if path == "/api/polish":
            qid = str(data.get("qid", "")).strip()
            query = str(data.get("query", "")).strip()
            if not qid:
                self._json(400, {"error": "안내 번호가 없습니다"})
                return
            started = time.perf_counter()
            try:
                out = polish(qid, query)
            except Exception as e:
                self._json(200, {"answer": "", "error": str(e)})
                return
            print(f"[다듬기] {time.perf_counter() - started:.1f}초 · {qid}", flush=True)
            self._json(200, out)
            return

        if path == "/api/guide":
            qid = str(data.get("qid", "")).strip()
            user = str(data.get("user", "")).strip()
            query = str(data.get("query", "")).strip()
            if not qid:
                self._json(400, {"error": "안내 번호가 없습니다"})
                return
            result = guide(qid, query)
            append_query_log(user, query or result.get("question") or qid, result, "안내선택")
            self._json(200, result)
            return

        try:
            if path == "/api/admin/save":
                self._json(200, save_item(data))
                return
            if path == "/api/admin/upload-form":
                self._json(200, upload_form(data))
                return
            if path == "/api/admin/upload-capture":
                self._json(200, upload_capture(data))
                return
            if path == "/api/admin/delete-capture":
                self._json(200, delete_capture(data))
                return
        except ValueError as e:
            self._json(400, {"error": str(e)})
            return
        except Exception as e:
            self._json(500, {"error": str(e)})
            return

        self._send(404, b"not found", "text/plain")


def main() -> None:
    global INDEX, META, BY_QID
    FORMS_DIR.mkdir(exist_ok=True)
    CAPTURES_DIR.mkdir(exist_ok=True)
    UPLOADS_DIR.mkdir(exist_ok=True)
    init_chat()
    INDEX = load_index()
    META = load_meta()
    BY_QID = {r["qid"]: r for r in INDEX}
    load_rewrite_cache()
    threading.Thread(target=keep_warm, daemon=True).start()
    print(f"색인 {len(INDEX)}건 · 답변 등록 {sum(1 for m in META.values() if m.get('answer'))}건", flush=True)
    print(f"직원 http://{HOST}:{PORT}", flush=True)
    print(f"관리 http://{HOST}:{PORT}/admin", flush=True)
    httpd = ThreadingHTTPServer((HOST, PORT), Handler)
    httpd.serve_forever()


if __name__ == "__main__":
    main()
